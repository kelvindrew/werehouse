#!/usr/bin/env python3
"""
Warehouse Management System (WMS) — B1 & B2 Excel Importer & Normalizer
Parses 'Copie de B1&B2 warehouse sheet-202602.xlsx', validates, cleans,
and generates normalized Firestore datasets with a full audit log.
"""

import openpyxl
import json
import os
import re
from datetime import datetime, timezone

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "..", "..", "drew.xlsx")
if not os.path.exists(EXCEL_FILE):
    EXCEL_FILE = r"C:\Users\Landry\Downloads\Copie de B1&B2 warehouse sheet-202602.xlsx"
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def safe_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0") and re.match(r"^\d+\.0$", s):
        s = s[:-2]
    return s

def safe_float(val, default=0.0):
    if val is None:
        return default
    s = str(val).strip()
    if s == "" or s.startswith("#"): # #N/A, #VALUE!, etc.
        return default
    try:
        return float(s)
    except:
        return default

def clean_bin(bin_val, warehouse_code):
    if not bin_val:
        return f"{warehouse_code}-UNASSIGNED"
    s = str(bin_val).strip()
    # Correct human typo: MDO1 -> MD01 (capital letter O instead of zero)
    if s.startswith("MDO1"):
        s = "MD01" + s[4:]
    return s

def run_import():
    print(f"Loading Excel file: {EXCEL_FILE} ...")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    timestamp = datetime.now(timezone.utc).isoformat()
    
    warehouses = [
        {
            "id": "B1",
            "code": "B1",
            "name": "Warehouse B1",
            "description": "Entrepôt principal B1 (Zone MD01)",
            "location": "Site Industriel - Bâtiment B1",
            "status": "ACTIVE",
            "createdAt": timestamp
        },
        {
            "id": "B2",
            "code": "B2",
            "name": "Warehouse B2",
            "description": "Entrepôt secondaire B2 (Zones B2-01 à B2-04, A, C, E)",
            "location": "Site Industriel - Bâtiment B2",
            "status": "ACTIVE",
            "createdAt": timestamp
        }
    ]
    
    # 1. Parse Sheet1 (Material Master Reference)
    sheet1_materials = {}
    if "Sheet1" in wb.sheetnames:
        ws1 = wb["Sheet1"]
        print(f"Parsing Sheet1 (Material Master), max_row={ws1.max_row}...")
        for r in range(2, ws1.max_row + 1):
            mat_raw = ws1.cell(row=r, column=1).value
            mat_code = safe_str(mat_raw)
            if not mat_code:
                continue
            sheet1_materials[mat_code] = {
                "plant": safe_str(ws1.cell(row=r, column=2).value),
                "valuationType": safe_str(ws1.cell(row=r, column=3).value),
                "description": safe_str(ws1.cell(row=r, column=4).value),
                "baseUom": safe_str(ws1.cell(row=r, column=6).value) or "EA",
                "price": safe_float(ws1.cell(row=r, column=7).value),
                "currency": safe_str(ws1.cell(row=r, column=8).value) or "USD",
                "priceUnit": safe_float(ws1.cell(row=r, column=9).value, default=1.0),
                "createdBy": safe_str(ws1.cell(row=r, column=10).value),
            }
        print(f"Loaded {len(sheet1_materials)} master material reference records from Sheet1.")

    materials = {} # materialCode -> dict
    bins = {} # binId -> dict
    stock = {} # stockId -> dict
    movements = []
    issues = []
    
    # Helper to register or update material
    def register_material(code, name, chinese_name, spec, uom, price, sheet1_info, requires_review=False):
        if code not in materials:
            base_price = price
            currency = "USD"
            plant = ""
            valuation_type = ""
            if sheet1_info:
                if base_price == 0.0 and sheet1_info["price"] > 0:
                    base_price = sheet1_info["price"]
                plant = sheet1_info["plant"]
                valuation_type = sheet1_info["valuationType"]
                currency = sheet1_info["currency"]
                
            materials[code] = {
                "id": code,
                "materialCode": code,
                "name": name or (sheet1_info["description"] if sheet1_info else "") or "Item without description",
                "chineseName": chinese_name or "",
                "description": name or "",
                "specification": spec or "",
                "uom": (uom or (sheet1_info["baseUom"] if sheet1_info else "") or "EA").upper(),
                "plant": plant,
                "valuationType": valuation_type,
                "standardPrice": base_price,
                "currency": currency,
                "priceUnit": 1.0,
                "requiresMaterialCodeReview": requires_review,
                "createdAt": timestamp,
                "updatedAt": timestamp
            }
        else:
            # Enrich existing material with any missing fields
            m = materials[code]
            if not m["chineseName"] and chinese_name:
                m["chineseName"] = chinese_name
            if not m["specification"] and spec:
                m["specification"] = spec
            if m["standardPrice"] == 0.0 and price > 0:
                m["standardPrice"] = price

    # 2. Parse B1 Warehouse Stock
    b1_summary = {
        "sheetName": "B1 warehouse stock",
        "totalRowsRead": 0,
        "validRows": 0,
        "phantomRowsSkipped": 0,
        "missingMaterialCodes": 0,
        "missingBins": 0,
        "missingPrices": 0,
        "duplicatesConsolidated": 0,
        "binTyposFixed": 0,
        "formulaErrorsResolved": 0,
        "originalSumQuantity": 0.0,
        "importedSumQuantity": 0.0,
        "originalSumTotalValue": 0.0,
        "importedSumTotalValue": 0.0,
        "isBalanced": False,
        "issues": []
    }
    
    ws_b1 = wb["B1 warehouse stock"]
    print(f"Parsing B1 warehouse stock, max_row={ws_b1.max_row}...")
    
    for r in range(2, ws_b1.max_row + 1):
        row_vals = [ws_b1.cell(row=r, column=c).value for c in range(1, 12)]
        # Check if entire row is empty
        if not any(v is not None for v in row_vals):
            continue
            
        b1_summary["totalRowsRead"] += 1
        
        c_no = safe_str(row_vals[0])
        c_mat = safe_str(row_vals[1])
        c_chinese = safe_str(row_vals[2])
        c_desc = safe_str(row_vals[3])
        c_spec = safe_str(row_vals[4])
        c_bin_raw = safe_str(row_vals[5])
        c_unit_price_raw = row_vals[6]
        c_qty_raw = row_vals[7]
        c_uom = safe_str(row_vals[8])
        c_total_price_raw = row_vals[9]
        c_remark = safe_str(row_vals[10])
        
        # Check for phantom row (only c_no exists, all others empty)
        data_cells = [c_mat, c_chinese, c_desc, c_spec, c_bin_raw, c_unit_price_raw, c_qty_raw, c_total_price_raw]
        if not any(v not in (None, "") for v in data_cells):
            b1_summary["phantomRowsSkipped"] += 1
            continue
            
        qty = safe_float(c_qty_raw)
        orig_total = safe_float(c_total_price_raw)
        unit_price = safe_float(c_unit_price_raw)
        b1_summary["originalSumQuantity"] += qty
        b1_summary["originalSumTotalValue"] += orig_total
        
        # Check typo in bin
        if c_bin_raw.startswith("MDO1"):
            b1_summary["binTyposFixed"] += 1
            b1_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B1",
                "type": "TYPO_CORRECTED",
                "field": "BIN Loc",
                "value": c_bin_raw,
                "message": "Typo MDO1 corrected to MD01",
                "actionTaken": "Replaced 'O' with '0'"
            })
            
        bin_loc = clean_bin(c_bin_raw, "B1")
        if not c_bin_raw:
            b1_summary["missingBins"] += 1
            b1_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B1",
                "type": "WARNING",
                "field": "BIN Loc",
                "value": None,
                "message": "Missing BIN location in B1",
                "actionTaken": "Assigned to B1-UNASSIGNED"
            })
            
        # Check material code
        requires_review = False
        if not c_mat:
            b1_summary["missingMaterialCodes"] += 1
            c_mat = f"TEMP-B1-R{r:04d}"
            requires_review = True
            b1_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B1",
                "type": "WARNING",
                "field": "Material",
                "value": None,
                "message": f"Missing Material Code for item '{c_desc}'",
                "actionTaken": f"Assigned temporary code {c_mat}"
            })
            
        # Price resolution
        s1_info = sheet1_materials.get(c_mat)
        if unit_price == 0.0:
            b1_summary["missingPrices"] += 1
            if s1_info and s1_info["price"] > 0:
                unit_price = s1_info["price"]
                
        # Register material
        register_material(c_mat, c_desc, c_chinese, c_spec, c_uom, unit_price, s1_info, requires_review)
        
        # Register BIN
        bin_id = f"B1_{bin_loc}"
        if bin_id not in bins:
            bins[bin_id] = {
                "id": bin_id,
                "warehouseId": "B1",
                "code": bin_loc,
                "zone": bin_loc.split("-")[0] if "-" in bin_loc else bin_loc,
                "status": "ACTIVE",
                "createdAt": timestamp
            }
            
        # Register Stock & Consolidate duplicates if same (warehouse, material, bin)
        stock_id = f"B1_{c_mat}_{bin_loc}"
        if stock_id in stock:
            b1_summary["duplicatesConsolidated"] += 1
            stock[stock_id]["quantity"] += qty
            stock[stock_id]["availableQuantity"] += qty
            stock[stock_id]["totalValue"] = stock[stock_id]["quantity"] * stock[stock_id]["unitPrice"]
            b1_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B1",
                "type": "DUPLICATE_CONSOLIDATED",
                "field": "Material + BIN",
                "value": f"{c_mat} @ {bin_loc}",
                "message": f"Duplicate item consolidated: added {qty} {c_uom}",
                "actionTaken": "Consolidated into existing stock record"
            })
        else:
            stock[stock_id] = {
                "id": stock_id,
                "materialId": c_mat,
                "materialCode": c_mat,
                "materialName": c_desc or c_mat,
                "chineseName": c_chinese,
                "specification": c_spec,
                "warehouseId": "B1",
                "binLocation": bin_loc,
                "uom": (c_uom or "EA").upper(),
                "quantity": qty,
                "reservedQuantity": 0.0,
                "availableQuantity": qty,
                "unitPrice": unit_price,
                "totalValue": round(qty * unit_price, 2),
                "remarks": c_remark,
                "lastUpdated": timestamp
            }
            
        # Record Stock Movement for initial import
        movements.append({
            "id": f"MOV-B1-{r:04d}",
            "movementType": "INITIAL_IMPORT",
            "materialId": c_mat,
            "materialCode": c_mat,
            "materialName": c_desc or c_mat,
            "warehouseId": "B1",
            "binLocation": bin_loc,
            "quantity": qty,
            "previousQuantity": 0.0,
            "newQuantity": qty,
            "unitPrice": unit_price,
            "totalAmount": round(qty * unit_price, 2),
            "referenceNumber": f"EXCEL-B1-ROW-{r}",
            "comments": f"Initial Excel import row {r}. Remark: {c_remark}".strip(),
            "performedBy": "SYSTEM_IMPORT",
            "performedByName": "Excel Import Migration",
            "createdAt": timestamp
        })
        
        b1_summary["validRows"] += 1

    # 3. Parse B2 Warehouse Stock
    b2_summary = {
        "sheetName": "B2 warehouse stock",
        "totalRowsRead": 0,
        "validRows": 0,
        "phantomRowsSkipped": 0,
        "missingMaterialCodes": 0,
        "missingBins": 0,
        "missingPrices": 0,
        "duplicatesConsolidated": 0,
        "binTyposFixed": 0,
        "formulaErrorsResolved": 0,
        "originalSumQuantity": 0.0,
        "importedSumQuantity": 0.0,
        "originalSumTotalValue": 0.0,
        "importedSumTotalValue": 0.0,
        "isBalanced": False,
        "issues": []
    }
    
    ws_b2 = wb["B2 warehouse stock"]
    print(f"Parsing B2 warehouse stock, max_row={ws_b2.max_row}...")
    
    # B2 header is at row 2, data starts row 3
    for r in range(3, ws_b2.max_row + 1):
        row_vals = [ws_b2.cell(row=r, column=c).value for c in range(1, 12)]
        if not any(v is not None for v in row_vals):
            continue
            
        b2_summary["totalRowsRead"] += 1
        
        c_no = safe_str(row_vals[0])
        c_bin_raw = safe_str(row_vals[1])
        c_mat = safe_str(row_vals[2])
        c_name = safe_str(row_vals[3])
        c_chinese = safe_str(row_vals[4])
        c_spec = safe_str(row_vals[5])
        c_qty_raw = row_vals[6]
        c_unit_price_raw = row_vals[7]
        c_amount_raw = row_vals[8]
        c_uom = safe_str(row_vals[9])
        c_remark = safe_str(row_vals[10])
        
        data_cells = [c_bin_raw, c_mat, c_name, c_chinese, c_spec, c_qty_raw, c_unit_price_raw, c_amount_raw]
        if not any(v not in (None, "") for v in data_cells):
            b2_summary["phantomRowsSkipped"] += 1
            continue
            
        qty = safe_float(c_qty_raw)
        orig_total = safe_float(c_amount_raw)
        unit_price = safe_float(c_unit_price_raw)
        
        if str(c_unit_price_raw or "").startswith("#") or str(c_amount_raw or "").startswith("#"):
            b2_summary["formulaErrorsResolved"] += 1
            b2_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B2",
                "type": "WARNING",
                "field": "Price/Amount",
                "value": f"Unit: {c_unit_price_raw}, Total: {c_amount_raw}",
                "message": "Excel formula error #N/A detected in price/amount cells",
                "actionTaken": "Sanitized to float 0.0 and checked Sheet1"
            })
            
        b2_summary["originalSumQuantity"] += qty
        b2_summary["originalSumTotalValue"] += orig_total
        
        bin_loc = clean_bin(c_bin_raw, "B2")
        if not c_bin_raw:
            b2_summary["missingBins"] += 1
            
        requires_review = False
        if not c_mat:
            b2_summary["missingMaterialCodes"] += 1
            c_mat = f"TEMP-B2-R{r:04d}"
            requires_review = True
            b2_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B2",
                "type": "WARNING",
                "field": "Material",
                "value": None,
                "message": f"Missing Material Code for item '{c_name}'",
                "actionTaken": f"Assigned temporary code {c_mat}"
            })
            
        s1_info = sheet1_materials.get(c_mat)
        if unit_price == 0.0:
            b2_summary["missingPrices"] += 1
            if s1_info and s1_info["price"] > 0:
                unit_price = s1_info["price"]
                
        register_material(c_mat, c_name, c_chinese, c_spec, c_uom, unit_price, s1_info, requires_review)
        
        bin_id = f"B2_{bin_loc}"
        if bin_id not in bins:
            bins[bin_id] = {
                "id": bin_id,
                "warehouseId": "B2",
                "code": bin_loc,
                "zone": bin_loc.split("-")[0] if "-" in bin_loc else bin_loc,
                "status": "ACTIVE",
                "createdAt": timestamp
            }
            
        stock_id = f"B2_{c_mat}_{bin_loc}"
        if stock_id in stock:
            b2_summary["duplicatesConsolidated"] += 1
            stock[stock_id]["quantity"] += qty
            stock[stock_id]["availableQuantity"] += qty
            stock[stock_id]["totalValue"] = stock[stock_id]["quantity"] * stock[stock_id]["unitPrice"]
            b2_summary["issues"].append({
                "rowNumber": r,
                "sheetName": "B2",
                "type": "DUPLICATE_CONSOLIDATED",
                "field": "Material + BIN",
                "value": f"{c_mat} @ {bin_loc}",
                "message": f"Duplicate item consolidated: added {qty} {c_uom}",
                "actionTaken": "Consolidated into existing stock record"
            })
        else:
            stock[stock_id] = {
                "id": stock_id,
                "materialId": c_mat,
                "materialCode": c_mat,
                "materialName": c_name or c_mat,
                "chineseName": c_chinese,
                "specification": c_spec,
                "warehouseId": "B2",
                "binLocation": bin_loc,
                "uom": (c_uom or "EA").upper(),
                "quantity": qty,
                "reservedQuantity": 0.0,
                "availableQuantity": qty,
                "unitPrice": unit_price,
                "totalValue": round(qty * unit_price, 2),
                "remarks": c_remark,
                "lastUpdated": timestamp
            }
            
        movements.append({
            "id": f"MOV-B2-{r:04d}",
            "movementType": "INITIAL_IMPORT",
            "materialId": c_mat,
            "materialCode": c_mat,
            "materialName": c_name or c_mat,
            "warehouseId": "B2",
            "binLocation": bin_loc,
            "quantity": qty,
            "previousQuantity": 0.0,
            "newQuantity": qty,
            "unitPrice": unit_price,
            "totalAmount": round(qty * unit_price, 2),
            "referenceNumber": f"EXCEL-B2-ROW-{r}",
            "comments": f"Initial Excel import row {r}. Remark: {c_remark}".strip(),
            "performedBy": "SYSTEM_IMPORT",
            "performedByName": "Excel Import Migration",
            "createdAt": timestamp
        })
        
        b2_summary["validRows"] += 1

    # Verify reconciliation sums
    for item in stock.values():
        if item["warehouseId"] == "B1":
            b1_summary["importedSumQuantity"] += item["quantity"]
            b1_summary["importedSumTotalValue"] += item["totalValue"]
        elif item["warehouseId"] == "B2":
            b2_summary["importedSumQuantity"] += item["quantity"]
            b2_summary["importedSumTotalValue"] += item["totalValue"]
            
    b1_summary["isBalanced"] = abs(b1_summary["importedSumQuantity"] - b1_summary["originalSumQuantity"]) < 0.001
    b2_summary["isBalanced"] = abs(b2_summary["importedSumQuantity"] - b2_summary["originalSumQuantity"]) < 0.001

    # Update warehouse totals
    warehouses[0]["totalItemsCount"] = sum(1 for s in stock.values() if s["warehouseId"] == "B1")
    warehouses[0]["totalQuantity"] = b1_summary["importedSumQuantity"]
    warehouses[0]["totalValuationUSD"] = round(b1_summary["importedSumTotalValue"], 2)
    
    warehouses[1]["totalItemsCount"] = sum(1 for s in stock.values() if s["warehouseId"] == "B2")
    warehouses[1]["totalQuantity"] = b2_summary["importedSumQuantity"]
    warehouses[1]["totalValuationUSD"] = round(b2_summary["importedSumTotalValue"], 2)

    # Save outputs as JSON files
    def save_json(filename, data):
        filepath = os.path.join(OUTPUT_DIR, filename)
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"Saved {filepath} ({len(data)} items)")

    save_json("warehouses.json", warehouses)
    save_json("materials.json", list(materials.values()))
    save_json("bins.json", list(bins.values()))
    save_json("stock.json", list(stock.values()))
    save_json("stock_movements.json", movements)
    
    audit_report = {
        "importTimestamp": timestamp,
        "b1Summary": b1_summary,
        "b2Summary": b2_summary,
        "totals": {
            "totalMaterials": len(materials),
            "totalBins": len(bins),
            "totalStockRecords": len(stock),
            "totalInitialMovements": len(movements),
            "overallQuantity": b1_summary["importedSumQuantity"] + b2_summary["importedSumQuantity"],
            "overallValuationUSD": round(b1_summary["importedSumTotalValue"] + b2_summary["importedSumTotalValue"], 2)
        }
    }
    save_json("import_audit_report.json", audit_report)
    
    print("\n================== IMPORT RECONCILIATION REPORT ==================")
    print(f"B1 Original QTY: {b1_summary['originalSumQuantity']} | Imported QTY: {b1_summary['importedSumQuantity']} | Balanced: {b1_summary['isBalanced']}")
    print(f"B1 Value: ${b1_summary['importedSumTotalValue']:,.2f} (Original: ${b1_summary['originalSumTotalValue']:,.2f})")
    print(f"B2 Original QTY: {b2_summary['originalSumQuantity']} | Imported QTY: {b2_summary['importedSumQuantity']} | Balanced: {b2_summary['isBalanced']}")
    print(f"B2 Value: ${b2_summary['importedSumTotalValue']:,.2f} (Original: ${b2_summary['originalSumTotalValue']:,.2f})")
    print("==================================================================")
    
    return audit_report

if __name__ == "__main__":
    run_import()
